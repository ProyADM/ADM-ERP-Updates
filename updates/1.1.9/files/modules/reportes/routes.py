# modules/reportes/routes.py
# ============================================================
# RUTAS DE REPORTES - SIDESYS ERP
# ============================================================
# Ahora solo delegan la lógica a los servicios.
# C4: cada endpoint exige permiso (reportes.ver/crear/eliminar/importar) y
# valida las bases pedidas (?base= / body / Excel) contra bases_permitidas.
# ============================================================

from flask import Blueprint, jsonify, request, session, send_file
import logging
import os
import re
import pandas as pd  # <--- AGREGAR
from datetime import datetime  # <--- AGREGAR
from .services import contratos_service, pendiente_cobro_service, consolidado_pais_service, consolidado_total_service, ventas_manuales_service
from modules.shared.decorators import (
    requiere_permiso,
    chequear_acceso_base,
    chequear_acceso_total_bases,
    chequear_acceso_alguna_base,
    bases_permitidas_actuales,
)
from modules.shared.permisos import PermisosSistema
from modules.shared import config_central
from modules.shared import correo_outlook

reportes_bp = Blueprint('reportes', __name__)
logger = logging.getLogger(__name__)

# ============================================================
# REPORTE: CONTRATOS PENDIENTE DE FACTURAR
# ============================================================

@reportes_bp.route('/reportes/contratos', methods=['POST'])
@requiere_permiso(PermisosSistema.REPORTES_VER)
def get_reporte_contratos():
    data = request.get_json() or {}
    base_override = data.get('base')
    if base_override:
        chk = chequear_acceso_base(base_override)
        if chk:
            return chk
    resultado = contratos_service.obtener_reporte_contratos(
        anio=data.get('anio'),
        mes=data.get('mes'),
        base_override=data.get('base'),
        sociedad_override=data.get('sociedad')
    )
    return jsonify(resultado)

# ============================================================
# REPORTE: AÑOS DISPONIBLES (para contratos)
# ============================================================

@reportes_bp.route('/reportes/anos', methods=['GET'])
@requiere_permiso(PermisosSistema.REPORTES_VER)
def get_anos_disponibles():
    # Esta función se mantiene igual, pero podemos moverla a utils si queremos
    # Por ahora la dejamos aquí porque es simple.
    try:
        from .services.utils import get_db_connection, get_db_connection_base
        base_param = request.args.get('base')
        sociedad_param = request.args.get('sociedad')
        if base_param:
            chk = chequear_acceso_base(base_param)
            if chk:
                return chk
        if base_param:
            conn, division = get_db_connection_base(base_param, sociedad_param)
        else:
            conn, division = get_db_connection()
        
        cursor = conn.cursor()
        cursor.execute("""
            SELECT DISTINCT YEAR(COFF_FECHA_EST_FC) AS ANIO
            FROM dbo.ACCT_COFF
            WHERE COFF_FECHA_EST_FC IS NOT NULL
              AND YEAR(COFF_FECHA_EST_FC) >= 2020
            ORDER BY ANIO DESC
        """)
        anos = [row[0] for row in cursor.fetchall()]
        conn.close()
        if not anos:
            from datetime import datetime
            ano_actual = datetime.now().year
            anos = [ano_actual, ano_actual - 1, ano_actual - 2, ano_actual - 3, ano_actual - 4]
        return jsonify({'success': True, 'anos': anos})
    except Exception as e:
        logger.error(f"Error en get_anos_disponibles: {e}")
        return jsonify({'success': False, 'error': str(e), 'anos': []}), 500

# ============================================================
# REPORTE: PENDIENTE DE COBRO
# ============================================================

@reportes_bp.route('/reportes/pendiente_cobro', methods=['GET'])
@requiere_permiso(PermisosSistema.REPORTES_VER)
def get_pendiente_cobro():
    base_override = request.args.get('base')
    sociedad_override = request.args.get('sociedad')
    # `fecha=AAAA-MM-DD` (opcional): el pendiente AL CORTE de esa fecha. Sin fecha (o con
    # la de hoy) sale el informe de siempre. Una fecha invalida se ignora (no rompe).
    fecha_corte = request.args.get('fecha')
    if base_override:
        chk = chequear_acceso_base(base_override)
        if chk:
            return chk
    resultado = pendiente_cobro_service.obtener_pendiente_cobro(
        base_override, sociedad_override, fecha_corte)
    return jsonify(resultado)

# ============================================================
# REPORTE: CONSOLIDADO VENTAS POR PAÍS (sin manuales)
# ============================================================

@reportes_bp.route('/reportes/consolidado_ventas_pais', methods=['GET'])
@requiere_permiso(PermisosSistema.REPORTES_VER)
def get_consolidado_ventas_pais():
    anio = request.args.get('anio')
    if anio and anio.isdigit():
        anio = int(anio)
    else:
        from datetime import datetime
        anio = datetime.now().year

    base_override = request.args.get('base')
    sociedad_override = request.args.get('sociedad')
    if base_override:
        chk = chequear_acceso_base(base_override)
        if chk:
            return chk
    # Sin ?base= este reporte recorre varias bases: se le pasa la lista del
    # usuario para que consulte SOLO las autorizadas (fail-closed).
    resultado = consolidado_pais_service.obtener_consolidado_pais(
        anio, base_override, sociedad_override,
        bases_permitidas=bases_permitidas_actuales())
    return jsonify(resultado)

# ============================================================
# REPORTE: CONSOLIDADO VENTAS GLOBAL (con manuales y cotización)
# ============================================================

@reportes_bp.route('/reportes/consolidado_ventas_global', methods=['GET'])
@requiere_permiso(PermisosSistema.REPORTES_VER)
def get_consolidado_ventas_global():
    # C4 (revisado 14/09/2026): antes exigía superadmin o bases '*'. Ahora alcanza
    # con tener AL MENOS UNA base, porque el reporte ya respeta `bases_permitidas`
    # (sólo consulta las bases del usuario) y devuelve `bases_incluidas` para que
    # la UI rotule que es un consolidado parcial. Un gerente con 2 bases puede ver
    # el consolidado DE SUS BASES sin que el resto de los países se filtre.
    chk = chequear_acceso_alguna_base()
    if chk:
        return chk

    anio = request.args.get('anio')
    if anio and anio.isdigit():
        anio = int(anio)
    else:
        from datetime import datetime
        anio = datetime.now().year

    base_override = request.args.get('base')
    sociedad_override = request.args.get('sociedad')
    if base_override:
        chk = chequear_acceso_base(base_override)
        if chk:
            return chk
    resultado = consolidado_total_service.obtener_consolidado_total(
        anio, base_override, sociedad_override,
        bases_permitidas=bases_permitidas_actuales())
    return jsonify(resultado)

# ============================================================
# REPORTE: AÑOS DISPONIBLES PARA VENTAS (para el selector del consolidado)
# ============================================================

@reportes_bp.route('/reportes/anos_ventas', methods=['GET'])
@requiere_permiso(PermisosSistema.REPORTES_VER)
def get_anos_ventas():
    try:
        from .services.utils import get_db_connection, get_db_connection_base
        base_override = request.args.get('base')
        sociedad_override = request.args.get('sociedad')
        if base_override:
            chk = chequear_acceso_base(base_override)
            if chk:
                return chk
        
        if base_override:
            conn, division = get_db_connection_base(base_override, sociedad_override)
        else:
            conn, division = get_db_connection()
        
        query = """
        SELECT DISTINCT YEAR(AASI_FECHA_REF) AS ANIO
        FROM SIST_AASI
        WHERE AASI_DIVISION = ?
          AND AASI_FECHA_REF IS NOT NULL
        ORDER BY ANIO DESC
        """
        df = pd.read_sql(query, conn, params=[division])
        conn.close()
        anos_bd = df['ANIO'].tolist() if not df.empty else []
        ano_actual = datetime.now().year
        anos_futuros = [ano_actual + 1, ano_actual, ano_actual - 1]
        todos_anos = list(set(anos_bd + anos_futuros))
        todos_anos.sort(reverse=True)
        anos_filtrados = [a for a in todos_anos if a >= 2025]
        if not anos_filtrados:
            anos_filtrados = [2027, 2026, 2025]
        anio_default = 2026
        return jsonify({'success': True, 'anos': anos_filtrados, 'anio_default': anio_default})
    except Exception as e:
        logger.error(f"Error en get_anos_ventas: {e}")
        return jsonify({'success': False, 'error': str(e), 'anos': [2027, 2026, 2025], 'anio_default': 2026}), 500

# ============================================================
# VENTAS MANUALES - CRUD
# ============================================================

@reportes_bp.route('/reportes/ventas_manuales', methods=['GET'])
@requiere_permiso(PermisosSistema.REPORTES_VER)
def get_ventas_manuales():
    base = request.args.get('base')
    anio = request.args.get('anio')
    mes = request.args.get('mes')
    if base:
        chk = chequear_acceso_base(base)
        if chk:
            return chk
    ventas = ventas_manuales_service.listar_ventas_manuales(base, anio, mes)

    # Alcance por país (17/09/2026). Con la tabla canónica las ventas de TODOS los
    # países viven juntas, así que hay que acotar acá:
    #   · el que puede con todo (superadmin / rol administrador) las ve todas;
    #   · el resto, solo las de sus países habilitados (fail-closed).
    # Cada fila lleva `puede_borrar` porque la decisión de borrado se calcula en
    # el servidor: así el frontend no puede "abrir" un borrado que el backend
    # igual rechazaría.
    from modules.shared.usuarios import get_gestor_usuarios
    username = session.get('username')
    gestor = get_gestor_usuarios()
    alcance_total = False
    try:
        datos = gestor.datos_usuario_completo(username) or {}
        alcance_total = bool(datos.get('es_superadmin')) or ('administrador' in (datos.get('roles') or []))
    except Exception as e:
        logger.warning(f'No se pudo resolver el alcance de ventas manuales: {e}')
        alcance_total = False
    permitidos = None if alcance_total else gestor.paises_permitidos(username)
    if permitidos is not None:
        ventas = [v for v in ventas if (v.get('pais') or '') in permitidos]
    for v in ventas:
        v['puede_borrar'] = bool(alcance_total or gestor.puede_gestionar_ventas_de(username, v.get('pais')))

    # Convertir fechas y decimales a tipos serializables
    for v in ventas:
        if v.get('fecha_registro'):
            v['fecha_registro'] = v['fecha_registro'].isoformat() if hasattr(v['fecha_registro'], 'isoformat') else str(v['fecha_registro'])
        if v.get('fecha_creacion'):
            v['fecha_creacion'] = v['fecha_creacion'].isoformat() if hasattr(v['fecha_creacion'], 'isoformat') else str(v['fecha_creacion'])
        if v.get('importe_usd'):
            v['importe_usd'] = float(v['importe_usd'])
        if v.get('importe_ps'):
            v['importe_ps'] = float(v['importe_ps'])
    return jsonify({'success': True, 'data': ventas, 'alcance_total': alcance_total})

@reportes_bp.route('/reportes/ventas_manuales', methods=['POST'])
@requiere_permiso(PermisosSistema.REPORTES_CREAR)
def crear_venta_manual():
    try:
        data = request.get_json()
        required = ['base', 'anio', 'mes', 'pais', 'fecha_registro']
        for field in required:
            if not data.get(field):
                return jsonify({'success': False, 'error': f'Falta campo: {field}'}), 400
        
        chk = chequear_acceso_base(data.get('base'))
        if chk:
            return chk

        importe_usd = data.get('importe_usd', 0)
        importe_ps = data.get('importe_ps', 0)
        if (importe_usd is None or importe_usd <= 0) and (importe_ps is None or importe_ps <= 0):
            return jsonify({'success': False, 'error': 'Debe ingresar al menos un importe (USD o PS)'}), 400
        
        resultado = ventas_manuales_service.crear_venta_manual(
            base=data['base'],
            sociedad=data.get('sociedad'),
            anio=int(data['anio']),
            mes=int(data['mes']),
            pais=data['pais'],
            importe_usd=importe_usd,
            importe_ps=importe_ps,
            fecha_registro=data['fecha_registro'],
            comentario=data.get('comentario', ''),
            usuario=session.get('username', 'anonimo'),
            centro_costo=data.get('centro_costo')
        )
        return jsonify(resultado)
    except Exception as e:
        logger.error(f"Error creando venta manual: {e}")
        return jsonify({'success': False, 'error': str(e)}), 500

@reportes_bp.route('/reportes/ventas_manuales/<int:venta_id>', methods=['DELETE'])
@requiere_permiso(PermisosSistema.REPORTES_ELIMINAR)
def eliminar_venta_manual(venta_id):
    """Borra una venta manual, validando el ALCANCE POR PAÍS.

    Regla (17/09/2026, pedido del usuario): el superadmin y el rol `administrador`
    pueden borrar las de cualquier país; el resto, solo las de sus países
    habilitados. Antes este endpoint borraba por id sin validar nada: alcanzaba
    con tener `reportes.eliminar` para borrar la venta de cualquier país — y con
    la tabla canónica eso quedaba más expuesto, porque las ventas de todos los
    países viven juntas.
    """
    from modules.shared.usuarios import get_gestor_usuarios

    # Para borrar, el chequeo previo se hace en modo ESCRITURA: si el usuario no
    # alcanza la base canónica, el error de permiso se propaga y se devuelve como
    # JSON legible en vez de convertirse en un 404 "La venta no existe" (que era
    # mentira: la venta existe, lo que falta es acceso).
    #
    # 200 + `success: false`, igual que los otros tres caminos del módulo (crear,
    # importar y el `jsonify(resultado)` del propio borrado): es un error de
    # NEGOCIO, y un 500 se lee como caída del servidor (dispara alertas y
    # monitoreo). (Ojo: el frontend del borrado hace `r.json()` sin mirar `r.ok`,
    # así que con 500 también vería el mensaje; el motivo es la consistencia.)
    try:
        venta = ventas_manuales_service.obtener_venta_manual(venta_id,
                                                            para_escritura=True)
    except Exception as e:
        logger.error(f'Error verificando la venta manual {venta_id} para borrar: {e}')
        return jsonify({'success': False, 'error': str(e)})

    if not venta:
        return jsonify({'success': False, 'error': 'La venta no existe'}), 404

    username = session.get('username')
    gestor = get_gestor_usuarios()
    if not gestor.puede_gestionar_ventas_de(username, venta.get('pais')):
        logger.warning(f'Borrado de venta manual denegado: {username} -> '
                       f'id={venta_id} pais={venta.get("pais")}')
        return jsonify({
            'success': False,
            'error': 'No podés borrar ventas de un país que no tenés habilitado.'
        }), 403

    resultado = ventas_manuales_service.eliminar_venta_manual(venta_id)
    if resultado.get('success'):
        logger.info(f'Venta manual {venta_id} eliminada por {username} '
                    f'(pais={venta.get("pais")})')
    return jsonify(resultado)

@reportes_bp.route('/reportes/ventas_manuales/importar', methods=['POST'])
@requiere_permiso(PermisosSistema.REPORTES_IMPORTAR)
def importar_ventas_manuales():
    try:
        if 'archivo' not in request.files:
            return jsonify({'success': False, 'error': 'No se envió archivo'}), 400
        
        archivo = request.files['archivo']
        if archivo.filename == '':
            return jsonify({'success': False, 'error': 'Nombre vacío'}), 400

        # Leer archivo con pandas
        try:
            if archivo.filename.endswith('.csv'):
                df = pd.read_csv(archivo, encoding='utf-8-sig')
            else:
                df = pd.read_excel(archivo)
        except Exception as e:
            return jsonify({'success': False, 'error': f'Error al leer el archivo: {str(e)}'}), 400

        # C4: validar que las bases indicadas en el Excel estén autorizadas.
        if 'base' in df.columns:
            bases_archivo = [str(b).strip() for b in df['base'].dropna().unique()]
            for b in bases_archivo:
                chk = chequear_acceso_base(b)
                if chk:
                    return chk

        resultado = ventas_manuales_service.importar_ventas_manuales(df, usuario=session.get('username', 'anonimo'))
        return jsonify(resultado)
    except Exception as e:
        logger.error(f"Error importando ventas manuales: {e}")
        return jsonify({'success': False, 'error': str(e)}), 500


# ============================================================
# EXPORTAR A .XLSX (reemplaza las descargas CSV del cliente)
# ============================================================
# Generico: el cliente envia {archivo, hoja, columnas, filas} y el servidor
# devuelve un .xlsx real (openpyxl). Toda exportacion de reportes pasa aca.
#
# Parametros OPCIONALES de presentacion (brief 23/09/2026, "Ventas Globales igual
# a la pantalla"). Si el cliente no los manda, NADA de esto se aplica y el
# archivo sale igual que antes:
#   formatos           lista alineada con `columnas`: numero de formato de Excel
#                      por columna ('' o null = celda como hoy).
#   anchos             lista de anchos de columna en caracteres; los indices con
#                      valor valido pisan el ancho automatico (el resto queda igual).
#   filas_negrita      indices 0-based de las filas de DATOS que van en negrita.
#   congelar_encabezado bool: true congela en la fila siguiente al encabezado,
#                      false lo suelta (si no viene, queda como hoy: congelado).
#   rellenos           lista de {'fila', 'columna', 'color'} (0-based, fila de
#                      datos y columna) con el color de fondo de esa celda.
#   sin_cuadricula     bool: true APAGA las lineas de cuadricula de la hoja
#                      (`ws.sheet_view.showGridLines = False`). Ausente (o
#                      cualquier otro valor) = como hoy, con cuadricula.
#   agrupaciones       lista de {'desde', 'hasta', 'colapsado'} con indices
#                      0-BASED de filas de DATOS (mismo criterio que
#                      `filas_negrita` y `rellenos`: la fila de encabezado es la
#                      1, asi que la fila de Excel es `indice + 2`). Aplica el
#                      agrupamiento NATIVO de Excel
#                      (`ws.row_dimensions.group(...)`, `outline_level=1`) y, si
#                      `colapsado` es verdadero, deja las filas OCULTAS. Ademas
#                      pone `ws.sheet_properties.outlinePr.summaryBelow = False`
#                      para que el `+/-` del grupo quede en la fila de ARRIBA (la
#                      del pais) y no abajo. Lo usa Ventas Globales para las
#                      filas de sociedad de Argentina (brief 24/09/2026). Un
#                      grupo invalido (`desde`/`hasta` no enteros, fuera de
#                      rango, `desde > hasta`) se IGNORA; ausente = como hoy, sin
#                      ningun grupo. Tope: 200 grupos por pedido.
#   autofiltro         bool: true pone el AUTOFILTRO del Excel
#                      (`ws.auto_filter.ref`) sobre TODA la tabla (de A1 a la
#                      ultima columna y la ultima fila, contando el encabezado).
#                      Ausente o cualquier otro valor = como hoy, sin autofiltro
#                      (y openpyxl NO escribe el `<autoFilter>`: el XML del libro
#                      sale igual al de antes). Lo piden los informes que tienen
#                      una columna por la que el usuario filtra en Excel
#                      (`Rango_Dias` en Pendiente de Cobro: el "segmentador
#                      RANGO DIAS DEUDA" del Excel de referencia, resuelto como
#                      columna + autofiltro).
#   hojas              lista de hojas para armar UN libro de VARIAS hojas:
#                      [{'hoja', 'columnas', 'filas', 'formatos', 'anchos',
#                      'filas_negrita', 'rellenos', 'agrupaciones',
#                      'agrupaciones_columnas', 'congelar_encabezado',
#                      'sin_cuadricula', 'autofiltro'}, ...]. Cada elemento tiene
#                      los MISMOS campos de arriba (los de presentacion son
#                      opcionales) y se arma con el MISMO codigo que la hoja
#                      unica: una sola implementacion, ningun reporte con su
#                      propio armado. Con `hojas` presente, los campos de
#                      presentacion del cuerpo se IGNORAN (cada hoja trae los
#                      suyos) y solo se usa `archivo` para el nombre del libro.
#                      Ausente, lista vacia, no-lista o con mas de 50 elementos =
#                      como hoy: UNA sola hoja con los campos del cuerpo (asi
#                      ningun informe existente cambia). La usa Ventas por Pais
#                      (Cliente, HW y Resto: tres hojas en un archivo).
#   formulas           matriz ALINEADA con `filas` (mismo alto y ancho) donde
#                      cada celda es `null` (queda el valor de `filas`) o el
#                      TEXTO de una formula de Excel (`'=SUM(E6:K6)'`). Se
#                      aplica AL FINAL, despues de formatos y estilos (asignar
#                      `.value` reemplaza la celda y se lleva el formato
#                      puesto). OJO: el archivo no lleva el valor cacheado de la
#                      formula; quien lo lea sin Excel (pandas, openpyxl con
#                      `data_only=True`) ve `None` hasta que Excel lo abra y lo
#                      guarde. Lo usa Pendiente de Facturar (la columna del anio
#                      y el Total general). Ausente = como hoy.
#   meses_visibles     rotulos de columna que el informe pide OCULTAR (no
#                      borrar): se marcan `hidden` para que sigan existiendo, y
#                      sumando en las formulas, pero no se vean. Se comparan
#                      contra el ENCABEZADO (fila 1) sin espacios de sobra y sin
#                      distinguir mayusculas (el motor no conoce los nombres de
#                      los meses: los manda el cliente). Lo usa Pendiente de
#                      Facturar para los meses cuyo total da cero. Ausente =
#                      como hoy (nada oculto).
#   agrupaciones_columnas
#                      lista de {'desde', 'hasta', 'colapsado'} con indices
#                      0-BASED de COLUMNAS de datos (mismo criterio que
#                      `rellenos`: la letra de Excel es `get_column_letter(indice
#                      + 1)`, o sea que el indice 0 es la columna A). Es el
#                      hermano de `agrupaciones` pero en las COLUMNAS: aplica
#                      `ws.column_dimensions.group(...)` con `outline_level=1` y,
#                      si `colapsado` es verdadero, deja las columnas OCULTAS.
#                      Ademas pone `ws.sheet_properties.outlinePr.summaryRight =
#                      False` (UNA sola vez, y solo si hay algun grupo valido)
#                      para que el `+/-` del grupo quede en la columna de la
#                      IZQUIERDA (el rotulo del anio). Lo usa Pendiente de
#                      Facturar para los meses de cada anio (brief 24/09/2026). Un
#                      grupo invalido (`desde`/`hasta` no enteros, fuera de
#                      rango, `desde > hasta`) se IGNORA; ausente = como hoy: sin
#                      ningun grupo y sin MODIFICAR `outlinePr` (este codigo no lo
#                      toca). OJO: openpyxl lo escribe SIEMPRE, con los defaults de
#                      Excel (`<outlinePr summaryBelow="1" summaryRight="1"/>`),
#                      que ya estaban antes de este cambio; por eso el XML del
#                      libro sale byte a byte igual al de antes (medido, ver
#                      `report.md`). Tope: 200 grupos.
#   resumen_columnas   de que lado del grupo de COLUMNAS queda la celda del `+/-`:
#                      `'izquierda'` (el default, como hoy: la columna del anio va
#                      DELANTE de sus meses) escribe `summaryRight = False`;
#                      `'derecha'` (la columna del anio va DESPUES de sus meses,
#                      como el libro de referencia del usuario) lo deja en `True`.
#                      Solo tiene efecto si hay `agrupaciones_columnas` validas;
#                      cualquier otro valor = izquierda.
#   filas_ocultas      indices 0-BASED de las FILAS de datos que van PLEGADAS
#                      (`hidden` + `outline_level=1`). Es el equivalente por FILA
#                      de `agrupaciones_columnas`: el `+`/`-` de Excel queda en la
#                      fila visible de arriba (la del total del grupo). Un indice
#                      invalido (no entero, fuera de rango) se IGNORA. Lo usa el
#                      pivote de Pendiente de Cobro (los comprobantes de cada
#                      cliente, plegados bajo su total). Ausente = como hoy.
#   columnas_ocultas   rotulos del ENCABEZADO (fila 1) de las columnas que se
#                      OCULTAN (no se borran): misma idea y misma comparacion que
#                      `meses_visibles` (sin espacios de sobra, sin distinguir
#                      mayusculas). Existe porque `meses_visibles` dice "meses" y
#                      ahora tambien se ocultan columnas que no son meses (los
#                      importes de Cobro que dan cero). Ausente = como hoy.
#   tabla              `true` = FORMATO DE TABLA de Excel (el "Dar formato como
#                      tabla" del catalogo), con franjas de fila. Opcionales:
#                      `tabla_estilo` (por defecto `TableStyleMedium2`, el azul
#                      "Medio 2" de la captura del usuario) y `tabla_nombre` (por
#                      defecto se arma con el nombre de la hoja). Se OMITE si la
#                      hoja no tiene filas (Excel no acepta una tabla vacia) o si el
#                      ENCABEZADO tiene rotulos vacios o repetidos (Excel los exige
#                      unicos y no vacios: con un encabezado asi REPARA el archivo al
#                      abrirlo, que es lo que el usuario vio el 05/10/2026). Las
#                      columnas de la tabla se escriben con los MISMOS rotulos que el
#                      encabezado de la hoja. OJO: una tabla trae los desplegables de
#                      filtro en el encabezado y `TableStyleInfo` no los controla. Lo
#                      usan los cuatro informes.

def _exportar_lista_opcional(data, nombre):
    """Lista del body o [] si no vino / no es lista (valor invalido = ignorar)."""
    valor = data.get(nombre)
    if not isinstance(valor, list):
        return []
    return valor


def _exportar_entero(valor):
    """Int del valor o None si no es un entero valido (no revienta)."""
    try:
        if isinstance(valor, bool):
            return None
        return int(valor)
    except (TypeError, ValueError):
        return None


def _exportar_formato(valor, indice, total_columnas):
    """Formato de Excel de una columna, o None si no corresponde aplicarlo."""
    if indice >= total_columnas or indice >= len(valor):
        return None
    formato = valor[indice]
    if not isinstance(formato, str) or not formato.strip():
        return None
    return formato[:60]


def _exportar_ancho(valor):
    """Ancho de columna en caracteres, o None si el valor no sirve.

    None (el `null` del body) o un valor no numerico o fuera de rango NO pisan el
    ancho automatico: la celda/columna queda "como hoy".
    """
    try:
        ancho = float(valor)
    except (TypeError, ValueError):
        return None
    if ancho != ancho:  # NaN
        return None
    if ancho <= 0 or ancho > 100:
        return None
    return ancho


def _exportar_color(valor):
    """Color ARGB/RGB de un relleno, o None si no es un string usable."""
    if not isinstance(valor, str) or not valor.strip():
        return None
    return valor.strip().lstrip('#')[:8].upper()


def _exportar_archivo_y_hoja(data):
    """Nombre del archivo y de la hoja saneados (lo comparten descarga y envio)."""
    import re as _re
    archivo = _re.sub(r'[\\/:*?"<>|]', '_', str(data.get('archivo') or 'exportacion.xlsx'))[:80]
    if not archivo.lower().endswith('.xlsx'):
        archivo += '.xlsx'
    hoja = _re.sub(r'[\[\]:*?/\\]', '', str(data.get('hoja') or 'Hoja1'))[:31] or 'Hoja1'
    return archivo, hoja


def construir_xlsx(data):
    """El .xlsx del pedido, como bytes. UNA SOLA funcion arma el libro.

    La usan la DESCARGA (`exportar_xlsx`) y el ENVIO por correo
    (`enviar_xlsx`): el adjunto del correo tiene que ser identico al archivo que
    el usuario baja (brief 23/09/2026). Si el armado se duplicara, los dos
    archivos se separarian con el primer cambio.

    El libro lo arma `_construir_libro_xlsx` (una hoja con los campos del cuerpo,
    o una por cada elemento de `hojas` si ese campo vino; ver el comentario del
    bloque de arriba) y cada hoja la escribe `_exportar_hoja`: aca solo se lo
    pasa a bytes.

    Devuelve `(bytes_del_xlsx, archivo, hoja)` o `(None, mensaje_de_error, None)`
    si el payload no sirve.
    """
    from io import BytesIO

    wb, archivo = _construir_libro_xlsx(data)
    if wb is None:
        return None, archivo, None
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue(), archivo, wb.active.title


def _construir_libro_xlsx(data):
    """El libro (Workbook) del pedido, ya escrito en memoria.

    Devuelve `(wb, archivo)` o `(None, mensaje_de_error)`. Lo comparten la
    descarga (`exportar_xlsx`) y el envio por correo (`enviar_xlsx`), los dos a
    traves de `construir_xlsx`: es el UNICO armado del libro.
    """
    from openpyxl import Workbook

    archivo, hoja = _exportar_archivo_y_hoja(data)
    wb = Workbook()

    # `hojas`: el libro de VARIAS hojas (ver el comentario del bloque de
    # arriba). Ausente, vacia, no-lista o con mas de 50 elementos = como hoy:
    # UNA sola hoja con los campos del cuerpo (ningun informe cambia).
    hojas = data.get('hojas')
    if isinstance(hojas, list) and 0 < len(hojas) <= 50:
        preparadas = []
        nombres_usados = set()
        for indice, definicion in enumerate(hojas):
            if not isinstance(definicion, dict):
                preparadas.append(None)
                continue
            # La PRIMERA hoja sin nombre propio toma el del cuerpo (los informes
            # de hoy mandan `hoja`; asi `hojas: [una]` arma EXACTAMENTE el mismo
            # archivo que el cuerpo de una sola hoja). Una hoja posterior sin
            # nombre se numera. Un nombre repetido se desambigua: dos hojas no
            # pueden llamarse igual.
            propio = definicion.get('hoja')
            if not propio:
                propio = hoja if indice == 0 else 'Hoja%d' % (indice + 1)
            nombre = _exportar_archivo_y_hoja({'hoja': propio})[1]
            if nombre in nombres_usados:
                sufijo = 2
                while ('%s%d' % (nombre[:29], sufijo)) in nombres_usados:
                    sufijo += 1
                nombre = '%s%d' % (nombre[:29], sufijo)
            nombres_usados.add(nombre)
            columnas, filas = _exportar_tabla(definicion)
            if filas is None:
                return None, 'Filas invalidas o demasiadas'
            preparadas.append((nombre, columnas, filas, definicion))

        for indice, preparada in enumerate(preparadas):
            if preparada is None:
                continue
            nombre, columnas, filas, definicion = preparada
            # La PRIMERA hoja reusa la que crea `Workbook()` (asi el archivo sale
            # con una hoja activa, como el de una sola hoja); las demas se crean.
            ws = wb.active if indice == 0 else wb.create_sheet()
            ws.title = nombre
            _exportar_hoja(ws, columnas, filas, definicion)
        if not wb.sheetnames:
            _exportar_hoja(wb.active, [], [], {})
        return wb, archivo

    # Sin `hojas` (o invalido): UNA sola hoja con la presentacion del cuerpo.
    columnas, filas = _exportar_tabla(data)
    if filas is None:
        return None, 'Filas invalidas o demasiadas'
    _exportar_hoja(wb.active, columnas, filas, data)
    wb.active.title = hoja
    return wb, archivo


def _exportar_tabla(data):
    """`(columnas, filas)` saneadas del cuerpo (o de UN elemento de `hojas`).

    Misma validacion de siempre: `filas` no-lista o con mas de 200000 elementos
    devuelve `(columnas, None)`, para que el que llama responda el error de hoy.
    """
    columnas = data.get('columnas') or []
    if not isinstance(columnas, list):
        columnas = []
    columnas = [str(c)[:60] for c in columnas][:60]

    # `filas` AUSENTE (o en `None`) = sin filas, como siempre: el libro sale con
    # el encabezado solo (lo fija `test_body_vacio...`: el cuerpo `{}` tiene que
    # exportar). Pero un `filas` PRESENTE y que no sea lista (un texto, un
    # numero) sigue siendo el 400 de siempre: `or []` solo, convertiria el `'no'`
    # en "sin filas" y el endpoint devolveria un archivo vacio en vez del error.
    filas = data.get('filas')
    if filas is None:
        filas = []
    if not isinstance(filas, list) or len(filas) > 200000:
        return columnas, None

    limpias = []
    for fila in filas:
        if not isinstance(fila, (list, tuple)):
            continue
        limpia = []
        for valor in list(fila)[:60]:
            if valor is None or isinstance(valor, (int, float, bool, str)):
                limpia.append(valor)
            else:
                limpia.append(str(valor)[:1000])
        limpias.append(limpia)
    return columnas, limpias


def _nombres_de_columnas_de_tabla(columnas):
    """Los rotulos del encabezado como nombres de columna de la tabla de Excel.

    Excel exige que los nombres de una tabla sean **unicos** y **no vacios**, y que
    sean los MISMOS que el encabezado de la hoja. Devuelve la lista, o `None` si el
    encabezado no sirve (hay un rotulo vacio o uno repetido): en ese caso la hoja se
    queda SIN tabla a proposito.

    Por que: si no se le pasan las columnas, openpyxl las arma EL MISMO al escribir
    (`WorksheetWriter.write_tables` -> `Table._initialise_columns`: nombra cada
    columna con el valor de la celda del encabezado) pero **NO desduplica**. Un
    encabezado con un rotulo repetido -como tenia Pendiente de Facturar, que rotulaba
    los meses de TODOS los anios con el numero pelado (`8` en 2026 y en 2027)- deja la
    tabla con **dos columnas del mismo nombre**, y Excel lee eso como un archivo roto y
    lo REPARA al abrirlo (es el cartel que vio el usuario el 05/10/2026: `Registros
    reparados: Tabla de /xl/tables/table1.xml`).
    """
    if not columnas:
        return None
    nombres = []
    for columna in columnas:
        nombre = str(columna).strip()
        if not nombre:
            return None
        nombres.append(nombre)
    if len(set(nombres)) != len(nombres):
        return None
    return nombres


def _exportar_tabla_con_estilo(ws, columnas, filas, opciones):
    """Le da FORMATO DE TABLA de Excel a la hoja ya escrita (brief 05/10/2026).

    El usuario pidio, con la captura del catalogo "Dar formato como tabla", el
    estilo azul de la fila del medio ("Medio 2"): el nombre de openpyxl es
    `TableStyleMedium2` y es el defecto de `tabla_estilo`.

    `ws` tiene el encabezado en la fila 1 (asi lo escribe `_exportar_hoja`), asi que
    la tabla va de A1 al ancho de `columnas` y hasta la ultima fila escrita. Se
    OMITE cuando no hay filas de datos (Excel no acepta una tabla sin datos) y
    cuando el encabezado tiene rotulos vacios o repetidos (ver
    `_nombres_de_columnas_de_tabla`).

    OJO (defecto corregido el 05/10/2026): la tabla se escribe **solo si el encabezado
    tiene rotulos unicos y no vacios** (ver `_nombres_de_columnas_de_tabla`), y sus
    columnas se escriben con ESOS rotulos. Antes se agregaba la tabla siempre y
    openpyxl nombraba las columnas con el encabezado **sin desduplicar**: un encabezado
    repetido (los meses de TODOS los anios rotulados con el numero pelado) dejaba la
    tabla con dos columnas del mismo nombre y Excel **reparaba el archivo** al abrirlo.

    OJO (declarado): una tabla de Excel lleva los desplegables de filtro en el
    encabezado; `TableStyleInfo` no los controla (solo el estilo) y openpyxl solo pone
    el `<autoFilter>` cuando es el el que arma las columnas, asi que este codigo lo
    escribe a mano.
    """
    if not columnas or not filas:
        return
    nombres = _nombres_de_columnas_de_tabla(columnas)
    if nombres is None:
        return
    from openpyxl.worksheet.filters import AutoFilter
    from openpyxl.worksheet.table import Table, TableColumn, TableStyleInfo
    from openpyxl.utils import get_column_letter

    estilo = opciones.get('tabla_estilo')
    if not isinstance(estilo, str) or not estilo.strip():
        estilo = 'TableStyleMedium2'
    nombre = opciones.get('tabla_nombre')
    if not isinstance(nombre, str) or not nombre.strip():
        # Los nombres de tabla de Excel no llevan espacios ni empiezan con digito:
        # se arma con el nombre de la HOJA, que ya es unico en el libro.
        nombre = 'Tabla_' + ''.join(
            caracter if caracter.isalnum() else '_' for caracter in ws.title)
    referencia = 'A1:%s%d' % (get_column_letter(len(columnas)), len(filas) + 1)
    tabla = Table(displayName=nombre, ref=referencia,
                  tableColumns=[TableColumn(id=indice + 1, name=rotulo)
                                for indice, rotulo in enumerate(nombres)])
    # La tabla trae los desplegables de filtro del encabezado: openpyxl los pone
    # cuando es EL el que arma las columnas (`WorksheetWriter.write_tables` ->
    # `Table._initialise_columns`). Al pasarle las columnas nosotros, el `autoFilter`
    # hay que ponerlo a mano para que el archivo salga igual que antes.
    tabla.autoFilter = AutoFilter(ref=referencia)
    tabla.tableStyleInfo = TableStyleInfo(
        name=estilo, showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False)
    ws.add_table(tabla)


def _exportar_hoja(ws, columnas, filas, opciones):
    """Escribe UNA hoja del libro: la hoja unica y cada hoja de `hojas: [...]`.

    `opciones` es el cuerpo del pedido (o el elemento de `hojas`): de ahi salen
    los campos de presentacion, que siguen siendo TODOS opcionales (ausente =
    como hoy). `columnas` y `filas` ya vienen saneadas por `_exportar_tabla`.
    El `Workbook` lo pone el que llama: aca solo se escribe la hoja.
    """
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    if columnas:
        ws.append(columnas)
        for cell in ws[1]:
            cell.font = Font(bold=True)
    for fila in filas:
        ws.append(fila)

    # Ancho aproximado de columnas (primeras 200 filas)
    for i, col in enumerate(ws.columns, 1):
        maxlen = 8
        for cell in list(col)[:200]:
            v = cell.value
            if v is not None:
                maxlen = max(maxlen, min(len(str(v)), 60))
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = maxlen + 2

    # ---- Presentacion opcional (si no viene, todo queda como hoy) --------
    formatos = _exportar_lista_opcional(opciones, 'formatos')
    anchos = _exportar_lista_opcional(opciones, 'anchos')
    if formatos or anchos:
        for indice, fila in enumerate(filas):
            if not isinstance(fila, (list, tuple)):
                continue
            for columna in range(min(len(fila), len(columnas))):
                formato = _exportar_formato(formatos, columna, len(columnas))
                if formato:
                    letra = get_column_letter(columna + 1)
                    ws[f'{letra}{indice + 2}'].number_format = formato
    for indice, valor in enumerate(anchos):
        if indice >= len(columnas):
            continue
        ancho = _exportar_ancho(valor)
        if ancho is not None:
            letra = get_column_letter(indice + 1)
            ws.column_dimensions[letra].width = ancho

    for bruto in _exportar_lista_opcional(opciones, 'filas_negrita'):
        indice = _exportar_entero(bruto)
        if indice is None or indice < 0 or indice >= len(filas):
            continue
        for columna in range(1, len(columnas) + 1):
            ws.cell(row=indice + 2, column=columna).font = Font(bold=True)

    # `filas_ocultas` (brief 05/10/2026): indices 0-based de las filas de DATOS
    # que van PLEGADAS (ocultas, con su nivel de esquema). Es el equivalente por
    # FILA de `agrupaciones_columnas`: el `+`/`-` de Excel queda en la fila de
    # arriba (la del cliente, que va visible), asi que el nivel se marca
    # en la fila oculta y Excel dibuja el boton solo. Lo usa el pivote de
    # Pendiente de Cobro (los comprobantes de cada cliente, plegados) y las dos
    # hojas de Pendiente de Facturar (los contratos de cada cliente).
    # `summaryBelow = False`: es el MISMO ajuste que hace `agrupaciones` (el resumen
    # va ARRIBA del grupo, que es como se arman los dos pivotes). Sin el, el boton
    # quedaba en la fila de ABAJO del grupo (o sea en el cliente siguiente), porque
    # el default de openpyxl es `summaryBelow="1"`. Solo se escribe si hay alguna
    # fila valida: con el campo ausente o invalido el XML sale igual que antes.
    hay_filas_ocultas = False
    for bruto in _exportar_lista_opcional(opciones, 'filas_ocultas'):
        indice = _exportar_entero(bruto)
        if indice is None or indice < 0 or indice >= len(filas):
            continue
        hay_filas_ocultas = True
        ws.row_dimensions[indice + 2].hidden = True
        ws.row_dimensions[indice + 2].outline_level = 1
    if hay_filas_ocultas:
        ws.sheet_properties.outlinePr.summaryBelow = False

    for relleno in _exportar_lista_opcional(opciones, 'rellenos'):
        if not isinstance(relleno, dict):
            continue
        indice = _exportar_entero(relleno.get('fila'))
        columna = _exportar_entero(relleno.get('columna'))
        color = _exportar_color(relleno.get('color'))
        if indice is None or columna is None or color is None:
            continue
        if indice < 0 or indice >= len(filas) or columna < 0 or columna >= len(columnas):
            continue
        ws.cell(row=indice + 2, column=columna + 1).fill = PatternFill(
            fill_type='solid', start_color=color, end_color=color)

    # `congelar_encabezado` AUSENTE -> como hoy (congelado en A2). Presente
    # en false -> el cliente pide explicitamente NO congelar.
    if opciones.get('congelar_encabezado') is False:
        ws.freeze_panes = None
    else:
        ws.freeze_panes = 'A2'

    # `sin_cuadricula` AUSENTE -> como hoy (la hoja mantiene la cuadricula: el
    # XML sale SIN el atributo, que es el default visible de Excel). Presente en
    # `true` -> `showGridLines="0"` (Ventas Globales). Cualquier otro valor se
    # ignora: el campo es opcional y los demas reportes no lo mandan.
    if opciones.get('sin_cuadricula') is True:
        ws.sheet_view.showGridLines = False

    # `agrupaciones` (brief 24/09/2026): agrupamiento NATIVO de Excel. Los indices
    # son 0-based de filas de DATOS -> fila de Excel = indice + 2. Ausente o
    # invalido = como hoy (sin ningun grupo): ningun otro informe cambia. OJO con
    # `outlinePr`: openpyxl lo escribe SIEMPRE, con los defaults de Excel
    # (`summaryBelow="1" summaryRight="1"`), asi que con el campo ausente el XML
    # del libro sale BYTE A BYTE igual al de antes del cambio, pero el bloque
    # esta (medido: auxiliar 322 de la entrega
    # `2026-09-24-pendiente-facturar-columnas-por-anio`; el comentario anterior
    # afirmaba "sin outlinePr" y era inexacto). `summaryBelow = False` solo se
    # escribe si hay algun grupo valido.
    for grupo in _exportar_lista_opcional(opciones, 'agrupaciones')[:200]:
        if not isinstance(grupo, dict):
            continue
        desde = _exportar_entero(grupo.get('desde'))
        hasta = _exportar_entero(grupo.get('hasta'))
        if desde is None or hasta is None:
            continue
        if desde < 0 or desde > hasta or hasta >= len(filas):
            continue
        # El `+/-` del grupo queda en la fila de ARRIBA (la del pais, que es el
        # padre), no abajo.
        ws.sheet_properties.outlinePr.summaryBelow = False
        ws.row_dimensions.group(desde + 2, hasta + 2, outline_level=1,
                                hidden=bool(grupo.get('colapsado')))

    # `agrupaciones_columnas` (brief 24/09/2026): el mismo agrupamiento NATIVO de
    # Excel pero en las COLUMNAS. Los indices son 0-based de COLUMNAS de datos
    # (mismo criterio que `rellenos`: la letra de Excel es `indice + 1`). Ausente
    # o invalido = como hoy: sin ningun grupo y sin MODIFICAR `outlinePr` (este
    # codigo no lo toca; openpyxl lo escribe SIEMPRE con los defaults de Excel,
    # `summaryBelow="1" summaryRight="1"`, tambien cuando el campo no viene, asi
    # que el XML sale byte a byte igual al de antes del cambio): ningun otro
    # informe cambia. `summaryRight = False` se escribe UNA sola vez y solo si hay
    # algun grupo valido: deja el `+/-` en la columna de la IZQUIERDA del grupo
    # (en Pendiente de Facturar, la del rotulo del anio).
    grupos_columnas = []
    for grupo in _exportar_lista_opcional(opciones, 'agrupaciones_columnas')[:200]:
        if not isinstance(grupo, dict):
            continue
        desde = _exportar_entero(grupo.get('desde'))
        hasta = _exportar_entero(grupo.get('hasta'))
        if desde is None or hasta is None:
            continue
        if desde < 0 or desde > hasta or hasta >= len(columnas):
            continue
        grupos_columnas.append((desde, hasta, bool(grupo.get('colapsado'))))
    if grupos_columnas:
        # `resumen_columnas` (05/10/2026): de que lado del grupo queda la celda del
        # `+/-`. `'izquierda'` (el default, como hoy) = la columna resumen va DELANTE
        # de las agrupadas; `'derecha'` = va DESPUES (los meses de cada anio y, a su
        # derecha, la columna del anio: el formato del libro de referencia de
        # Pendiente de Facturar). Cualquier otro valor = izquierda.
        # `True` = el `+/-` va a la DERECHA del grupo (`'derecha'`); `False` = a la
        # izquierda (el default y cualquier otro valor).
        ws.sheet_properties.outlinePr.summaryRight = (
            opciones.get('resumen_columnas') == 'derecha')
        for desde, hasta, colapsado in grupos_columnas:
            ws.column_dimensions.group(get_column_letter(desde + 1),
                                       get_column_letter(hasta + 1),
                                       outline_level=1, hidden=colapsado)

    # `autofiltro` AUSENTE -> como hoy (sin autofiltro y sin escribir el
    # `<autoFilter>` del XML: openpyxl no lo emite cuando queda en `None`).
    # Presente en `true` -> la flechita de Excel sobre TODA la tabla (encabezado
    # incluido). Cualquier otro valor se ignora, como el resto de los campos
    # opcionales. Lo pide Pendiente de Cobro para su `Rango_Dias`.
    if opciones.get('autofiltro') is True and columnas and filas:
        ws.auto_filter.ref = 'A1:%s%d' % (get_column_letter(len(columnas)),
                                          len(filas) + 1)

    # `formulas` (brief 05/10/2026): matriz ALINEADA con `filas` (mismo alto y
    # ancho) donde cada celda es `null` (queda el valor) o el TEXTO de una
    # formula (`'=SUM(E6:K6)'`). Se aplica AL FINAL A PROPOSITO: asignar
    # `.value` en openpyxl reemplaza la celda y se lleva puesto el formato de
    # numero y el estilo, asi que primero van formatos/negritas/rellenos y
    # despues la formula. `data_only` en Excel: el archivo NO lleva el valor
    # cacheado, asi que quien lo lea con `openpyxl(data_only=True)` ve `None`
    # hasta que Excel abra y guarde el archivo; por eso el cliente puede mandar
    # tambien el valor en `filas` (la formula pisa el valor, no al reves).
    # Ausente, no-lista o con filas/columnas que no existen = como hoy.
    formulas = _exportar_lista_opcional(opciones, 'formulas')
    if formulas:
        celdas_con_formula = 0
        for indice, fila_formulas in enumerate(formulas[:len(filas)]):
            if not isinstance(fila_formulas, (list, tuple)):
                continue
            for columna, formula in enumerate(list(fila_formulas)[:len(columnas)]):
                if not isinstance(formula, str) or not formula.startswith('='):
                    continue
                ws.cell(row=indice + 2, column=columna + 1).value = formula[:500]
                celdas_con_formula += 1
        if celdas_con_formula:
            logger.debug(f"Exportacion: {celdas_con_formula} celdas con formula")

    # `meses_visibles` (brief 05/10/2026): rotulos de las columnas que el
    # informe decide OCULTAR (no borrar). Se ocultan para que sigan existiendo
    # (y sumando en las formulas), pero no se vean: el caso medido son los
    # meses cuyo total da cero. Se comparan los rotulos del ENCABEZADO (fila 1)
    # normalizados (sin espacios de sobra, en mayusculas) porque el motor no
    # conoce los nombres de los meses: los manda el cliente. Una columna que no
    # aparece se ignora; ausente o no-lista = como hoy (nada oculto).
    rotulos_ocultos = {str(r).strip().upper()
                       for r in _exportar_lista_opcional(opciones, 'meses_visibles')
                       if isinstance(r, str) and str(r).strip()}
    if rotulos_ocultos and columnas:
        for posicion in range(len(columnas)):
            celda = ws.cell(row=1, column=posicion + 1)
            if celda.value is None:
                continue
            if str(celda.value).strip().upper() in rotulos_ocultos:
                ws.column_dimensions[get_column_letter(posicion + 1)].hidden = True

    # `columnas_ocultas` (brief 05/10/2026): lo MISMO que `meses_visibles` (ocultar
    # por ROTULO del encabezado, no por indice) pero con un nombre que dice lo que
    # hace. Lo usa el pivote de Pendiente de Cobro: "las columnas que no tienen
    # importe que no se muestren". Ausente = como hoy.
    rotulos_ocultos_extra = {str(r).strip().upper()
                             for r in _exportar_lista_opcional(opciones, 'columnas_ocultas')
                             if isinstance(r, str) and str(r).strip()}
    if rotulos_ocultos_extra and columnas:
        for posicion in range(len(columnas)):
            celda = ws.cell(row=1, column=posicion + 1)
            if celda.value is None:
                continue
            if str(celda.value).strip().upper() in rotulos_ocultos_extra:
                ws.column_dimensions[get_column_letter(posicion + 1)].hidden = True

    # El FORMATO DE TABLA, ULTIMO: va despues de formatos, negritas y formulas, asi
    # el estilo de la tabla es lo que define los colores (y no se pisan entre si).
    if opciones.get('tabla'):
        _exportar_tabla_con_estilo(ws, columnas, filas, opciones)


@reportes_bp.route('/reportes/exportar_xlsx', methods=['POST'])
@requiere_permiso(PermisosSistema.REPORTES_VER)
def exportar_xlsx():
    try:
        from io import BytesIO

        data = request.get_json(silent=True)
        if not isinstance(data, dict):
            return jsonify({'success': False, 'error': 'Cuerpo requerido'}), 400

        contenido, archivo, _hoja = construir_xlsx(data)
        if contenido is None:
            return jsonify({'success': False, 'error': archivo}), 400

        return send_file(
            BytesIO(contenido),
            as_attachment=True,
            download_name=archivo,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
    except Exception as e:
        logger.error(f"Error exportando xlsx: {e}")
        return jsonify({'success': False, 'error': 'Error generando el archivo'}), 500


# ============================================================
# ENVIAR EL .XLSX POR CORREO (envio directo con el Outlook del usuario)
# ============================================================
# Brief 23/09/2026 ("Enviar el informe por correo desde Outlook") y brief
# 24/09/2026 ("Argentina por sociedad, nombre del archivo y envio directo"),
# que revierte la decision anterior: el correo ya NO queda como borrador abierto,
# se ENVIA (`Send()`, ver `modules/shared/correo_outlook.py`).
#
# El cuerpo es el MISMO de `exportar_xlsx` mas `para`/`asunto`/`cuerpo`: el
# adjunto sale de `construir_xlsx`, asi que es identico al de la descarga.

# Carpeta del adjunto: Outlook referencia el archivo por ruta mientras lo manda,
# asi que no se borra al terminar el request (se van conservando los ultimos N;
# ver `_podar_correos`).
CARPETA_CORREOS = os.path.join(os.environ.get('PROGRAMDATA', r'C:\ProgramData'),
                               'SidesysERP', 'correos')
MAX_CORREOS = 20

# Una direccion con forma valida: algo@algo.algo, sin espacios ni separadores.
PATRON_CORREO = re.compile(r'^[A-Za-z0-9._%+\-]+@[A-Za-z0-9]([A-Za-z0-9\-]*[A-Za-z0-9])?'
                           r'(\.[A-Za-z0-9]([A-Za-z0-9\-]*[A-Za-z0-9])?)+$')

# `config_central.auditar` es el mecanismo de auditoria del proyecto (una linea
# JSON en el almacen central, con respaldo en data/). Se deja en una constante
# para que las pruebas puedan doblarlo.
AUDITORIA = config_central.auditar

# La funcion que ENVIA el informe. Tambien en constante para poder doblarla:
# NINGUNA prueba abre Outlook ni manda un correo real.
ENVIAR_CORREO = correo_outlook.enviar_informe


def destinatarios_validos(para):
    """`(validos, invalidos)` del campo `para`.

    Separadores: `;` y `,` (los dos, como pide el brief). Una direccion sin
    forma valida se descarta y se informa: el endpoint responde 400 si no quedo
    NINGUNA valida.
    """
    validos = []
    invalidos = []
    texto = '' if para is None else str(para)
    for bruto in re.split(r'[;,]', texto):
        direccion = bruto.strip()
        if not direccion:
            continue
        if PATRON_CORREO.match(direccion):
            validos.append(direccion)
        else:
            invalidos.append(direccion)
    return validos, invalidos


def _podar_correos(carpeta, maximo=MAX_CORREOS):
    """Borra los adjuntos mas viejos de `carpeta`. Devuelve cuantos borro.

    Solo mira los `.xlsx` (los adjuntos que escribe este endpoint): un archivo
    ajeno en la carpeta no se toca. Nunca levanta por el disco: la retencion no
    puede tumbar el envio.
    """
    try:
        nombres = [n for n in os.listdir(carpeta) if n.lower().endswith('.xlsx')]
    except OSError:
        return 0
    rutas = []
    for nombre in nombres:
        completa = os.path.join(carpeta, nombre)
        try:
            rutas.append((os.path.getmtime(completa), completa))
        except OSError:
            continue
    borrados = 0
    for _fecha, completa in sorted(rutas, reverse=True)[maximo:]:
        try:
            os.remove(completa)
            borrados += 1
        except OSError as e:
            logger.warning(f"No se pudo borrar el adjunto viejo {completa}: {e}")
    return borrados


@reportes_bp.route('/reportes/enviar_xlsx', methods=['POST'])
@requiere_permiso(PermisosSistema.REPORTES_VER)
def enviar_xlsx():
    data = request.get_json(silent=True)
    if not isinstance(data, dict):
        return jsonify({'success': False, 'detalle': 'Cuerpo requerido'}), 400

    # 1) Destinatarios: se valida ANTES de armar el Excel (no se gasta trabajo en
    #    un pedido que no puede salir).
    validos, invalidos = destinatarios_validos(data.get('para'))
    if not validos:
        detalle = ('Indicá al menos un destinatario con dirección válida '
                   '(varios separados por ";" o ",").')
        if invalidos:
            detalle += ' No se reconocieron: ' + ', '.join(invalidos[:5])
        return jsonify({'success': False, 'detalle': detalle}), 400

    asunto = str(data.get('asunto') or '').replace('\r', ' ').replace('\n', ' ').strip()
    # Los saltos de linea se planan y el espacio sobrante se colapsa: el asunto
    # de un correo es de UNA linea (nunca un encabezado inyectado).
    asunto = ' '.join(asunto.split())
    if len(asunto) > 300:
        return jsonify({'success': False, 'detalle':
                        'El asunto es demasiado largo (máximo 300 caracteres).'}), 400
    cuerpo = str(data.get('cuerpo') or '')

    # 2) El adjunto: EXACTAMENTE el mismo armado que la descarga.
    try:
        contenido, archivo, _hoja = construir_xlsx(data)
    except Exception as e:
        logger.error(f"Error armando el xlsx para enviar: {e}")
        return jsonify({'success': False, 'detalle': 'Error generando el archivo'}), 500
    if contenido is None:
        return jsonify({'success': False, 'detalle': archivo}), 400

    # 3) El archivo en disco: el borrador de Outlook lo referencia por ruta.
    try:
        os.makedirs(CARPETA_CORREOS, exist_ok=True)
        nombre = f'{datetime.now().strftime("%Y%m%d_%H%M%S_%f")}_{os.path.basename(archivo)}'
        ruta_adjunto = os.path.join(CARPETA_CORREOS, nombre)
        with open(ruta_adjunto, 'wb') as f:
            f.write(contenido)
    except Exception as e:
        logger.error(f"No se pudo guardar el adjunto del correo: {e}")
        return jsonify({'success': False, 'detalle':
                        'No se pudo guardar el archivo del informe para adjuntar.'}), 500
    _podar_correos(CARPETA_CORREOS)

    # 4) El envio por Outlook. `enviar_informe` no propaga excepciones.
    #    `nombre_adjunto`: el archivo en disco lleva el prefijo de fecha/hora (para la
    #    retencion y para que dos envios no se pisen) y el correo tiene que llevar el
    #    nombre del informe, el mismo que muestra el dialogo (defecto reportado el
    #    05/10/2026: al destinatario le llegaba `20261005_095747_431465_RD - ...xlsx`).
    ok, detalle = ENVIAR_CORREO(ruta_adjunto, '; '.join(validos),
                                asunto or os.path.basename(archivo), cuerpo,
                                nombre_adjunto=os.path.basename(archivo))

    # 5) Auditoria: quien, a quienes y que informe. `auditar` nunca lanza; el try
    #    es por si el mecanismo no estuviera disponible (nunca una excepcion
    #    despues de haber mandado el correo).
    try:
        AUDITORIA({'evento': 'reportes.enviar_xlsx',
                   'usuario': session.get('username') or 'desconocido',
                   'para': validos,
                   'informe': os.path.basename(archivo),
                   'asunto': asunto,
                   'filas': len(data.get('filas') or []),
                   'adjunto': ruta_adjunto,
                   'resultado': 'enviado' if ok else 'fallido',
                   'detalle': detalle})
    except Exception as e:
        logger.error(f"No se pudo auditar el envio del informe: {e}")

    return jsonify({'success': bool(ok), 'detalle': detalle})
