# modules/cotizaciones/routes.py
# ============================================================
# RUTAS DE COTIZACIONES - SIDESYS ERP (VERSIÓN MEJORADA)
# ============================================================

from datetime import date, datetime
import base64
import os
import tempfile
import logging
import zipfile
from flask import Blueprint, jsonify, request
import openpyxl
import pyodbc
from config import (
    BASES_COTI,
    SQL_SERVER,
    SQL_USERNAME,
    SQL_PASSWORD,
)

# Configurar logging
logger = logging.getLogger(__name__)

cotizaciones_bp = Blueprint("cotizaciones", __name__, url_prefix="/api/cotizaciones")

__all__ = ["cotizaciones_bp"]

# ============================================================
# PAÍSES PERMITIDOS (VALIDACIÓN)
# ============================================================

PAISES_PERMITIDOS = list(BASES_COTI.keys())

# Países prioritarios para obtener cotización (primero RD, luego el resto)
PAISES_PRIORITARIOS = ['RD'] + [p for p in PAISES_PERMITIDOS if p != 'RD']

def validar_sigla(sigla: str) -> bool:
    """Valida que la sigla del país sea válida"""
    return sigla in PAISES_PERMITIDOS

def validar_fecha(fecha: str) -> bool:
    """Valida que la fecha tenga formato YYYY-MM-DD"""
    try:
        datetime.strptime(fecha, '%Y-%m-%d')
        return True
    except (ValueError, TypeError):
        return False

def validar_cotizacion(valor) -> bool:
    """Valida que la cotización sea un número positivo"""
    try:
        val = float(valor)
        return val > 0
    except (ValueError, TypeError):
        return False

def sanitizar_fecha(fecha: str) -> str:
    """Sanitiza una fecha para evitar inyección"""
    if not validar_fecha(fecha):
        raise ValueError(f"Fecha inválida: {fecha}")
    return fecha

# ============================================================
# CONEXIÓN DIRECTA POR PAÍS
# ============================================================

def obtener_conexion_pais(sigla):
    """Obtiene conexión a la base de datos de un país específico"""
    if not validar_sigla(sigla):
        raise ValueError(f"País no válido: {sigla}")
    
    db_name = BASES_COTI.get(sigla)
    if not db_name:
        raise ValueError(f"No existe base de datos para la sigla: {sigla}")
    
    try:
        from config import SQL_SERVER_PLATAFORMA
    except ImportError:
        SQL_SERVER_PLATAFORMA = SQL_SERVER

    server = SQL_SERVER_PLATAFORMA if sigla == "AR" else SQL_SERVER
    
    conn_str = (
        'DRIVER={SQL Server};'
        f'SERVER={server};'
        f'DATABASE={db_name};'
        f'UID={SQL_USERNAME};'
        f'PWD={SQL_PASSWORD};'
        'TrustServerCertificate=yes;'
        'Timeout=10;'            # Timeout de conexión (segundos)
        'Connect Timeout=10;'    # Alternativa
    )
    return pyodbc.connect(conn_str)

def obtener_conexion_pais_segura(sigla):
    """Obtiene conexión con manejo de errores"""
    try:
        return obtener_conexion_pais(sigla)
    except pyodbc.Error as e:
        logger.error(f"Error de conexión a {sigla}: {e}")
        raise RuntimeError(f"No se pudo conectar a la base de datos de {sigla}")
    except Exception as e:
        logger.error(f"Error inesperado en {sigla}: {e}")
        raise

# ============================================================
# FUNCIÓN AUXILIAR PARA GUARDAR COTIZACIONES (EVITA DUPLICACIÓN)
# ============================================================

def _guardar_cotizaciones(items):
    """
    Guarda una lista de cotizaciones en las bases correspondientes.
    
    Args:
        items: Lista de diccionarios con keys: fecha, cotizacion, bases (lista de siglas)
    
    Returns:
        dict: {ok, insertados, errores, bases_con_error}
    """
    if not items:
        return {"ok": False, "insertados": 0, "errores": ["No hay items para guardar"], "bases_con_error": []}
    
    errores = []
    ok = 0
    bases_con_error = []
    
    for item in items:
        fecha = item["fecha"]
        cotizacion = item["cotizacion"]
        bases = item.get("bases", PAISES_PERMITIDOS)
        
        # Validar que las bases sean válidas
        bases_validas = [b for b in bases if b in PAISES_PERMITIDOS]
        if not bases_validas:
            errores.append(f"No hay bases válidas para la fecha {fecha}")
            continue
        
        for sigla in bases_validas:
            conn = None
            try:
                conn = obtener_conexion_pais_segura(sigla)
                cursor = conn.cursor()

                # Consulta UPSERT con parámetros
                query = """
                    IF EXISTS (SELECT 1 FROM SIST_COTI WHERE COTI_MONEDA1=? AND COTI_MONEDA2=? AND COTI_FECHA=?)
                        UPDATE SIST_COTI SET COTI_COTIZACION=?
                        WHERE COTI_MONEDA1=? AND COTI_MONEDA2=? AND COTI_FECHA=?
                    ELSE
                        INSERT INTO SIST_COTI (COTI_MONEDA1, COTI_MONEDA2, COTI_FECHA, COTI_COTIZACION)
                        VALUES (?, ?, ?, ?)
                """
                cursor.execute(query, (
                    'DL', 'PS', fecha,
                    cotizacion,
                    'DL', 'PS', fecha,
                    'DL', 'PS', fecha, cotizacion
                ))
                conn.commit()
                ok += 1
                logger.info(f"Cotización {cotizacion} guardada en {sigla} para fecha {fecha}")
                
            except pyodbc.Error as e:
                error_msg = f"{sigla} {fecha}: {str(e)}"
                errores.append(error_msg)
                bases_con_error.append(sigla)
                logger.error(f"Error SQL en {sigla}: {e}")
            except Exception as e:
                error_msg = f"{sigla} {fecha}: {str(e)}"
                errores.append(error_msg)
                bases_con_error.append(sigla)
                logger.error(f"Error inesperado en {sigla}: {e}")
            finally:
                if conn:
                    try:
                        conn.close()
                    except:
                        pass
    
    return {
        "ok": len(errores) == 0,
        "insertados": ok,
        "errores": errores,
        "bases_con_error": list(set(bases_con_error))
    }

# ============================================================
# RUTAS
# ============================================================

@cotizaciones_bp.route("/historico")
def historico():
    """Obtiene el historial de cotizaciones por país"""
    result = {}
    
    for sigla in PAISES_PERMITIDOS:
        conn = None
        try:
            conn = obtener_conexion_pais_segura(sigla)
            cursor = conn.cursor()
            
            cursor.execute("""
                SELECT 
                    CONVERT(VARCHAR, COTI_FECHA, 23) as fecha,
                    COTI_MONEDA1 as moneda1,
                    COTI_MONEDA2 as moneda2,
                    COTI_COTIZACION as cotizacion
                FROM SIST_COTI
                WHERE COTI_MONEDA1 = ? AND COTI_MONEDA2 = ?
                ORDER BY COTI_FECHA DESC
            """, ('DL', 'PS'))
            
            columns = [column[0] for column in cursor.description]
            rows = []
            for row in cursor.fetchall():
                row_dict = dict(zip(columns, row))
                if row_dict.get("cotizacion") is not None:
                    try:
                        row_dict["cotizacion"] = float(str(row_dict["cotizacion"]).replace(",", "."))
                    except (ValueError, TypeError):
                        row_dict["cotizacion"] = 0.0
                rows.append(row_dict)
                
            result[sigla] = rows
            
        except ValueError as e:
            logger.warning(f"País inválido: {sigla} - {e}")
            result[sigla] = {"error": f"País no válido: {sigla}"}
        except pyodbc.Error as e:
            logger.error(f"Error SQL en {sigla}: {e}")
            result[sigla] = {"error": f"Error de base de datos: {str(e)}"}
        except Exception as e:
            logger.error(f"Error inesperado en {sigla}: {e}")
            result[sigla] = {"error": f"Error inesperado: {str(e)}"}
        finally:
            if conn:
                try:
                    conn.close()
                except:
                    pass
                    
    return jsonify(result)


@cotizaciones_bp.route("/guardar", methods=["POST"])
def guardar():
    """Guarda una cotización en uno o varios países"""
    try:
        data = request.json
        
        if not data:
            return jsonify({
                "ok": False,
                "error": "No se recibieron datos"
            }), 400
        
        if not isinstance(data, list):
            return jsonify({
                "ok": False,
                "error": "Se esperaba una lista de cotizaciones"
            }), 400
        
        items = []
        for item in data:
            # Validar campos
            if not item.get("fecha"):
                return jsonify({"ok": False, "error": "Fecha no proporcionada"}), 400
                
            if not validar_fecha(item["fecha"]):
                return jsonify({"ok": False, "error": f"Fecha inválida: {item['fecha']}"}), 400
                
            try:
                cotizacion = float(item["cotizacion"])
                if cotizacion <= 0:
                    return jsonify({"ok": False, "error": f"Cotización inválida (debe ser positiva): {cotizacion}"}), 400
            except (ValueError, TypeError):
                return jsonify({"ok": False, "error": f"Cotización inválida: {item.get('cotizacion')}"}), 400

            bases = item.get("bases", PAISES_PERMITIDOS)
            bases_validas = [b for b in bases if b in PAISES_PERMITIDOS]
            if not bases_validas:
                return jsonify({"ok": False, "error": f"No hay bases válidas para: {bases}"}), 400
            
            items.append({
                "fecha": item["fecha"],
                "cotizacion": cotizacion,
                "bases": bases_validas
            })
        
        resultado = _guardar_cotizaciones(items)
        
        if resultado["bases_con_error"]:
            return jsonify({
                "ok": resultado["ok"],
                "insertados": resultado["insertados"],
                "errores": resultado["errores"],
                "bases_con_error": resultado["bases_con_error"],
                "mensaje": f"⚠️ No se pudo guardar en: {', '.join(resultado['bases_con_error'])}."
            })
        
        return jsonify({
            "ok": resultado["ok"],
            "insertados": resultado["insertados"],
            "errores": resultado["errores"]
        })
        
    except Exception as e:
        logger.error(f"Error en guardar cotización: {e}")
        return jsonify({
            "ok": False,
            "error": f"Error inesperado: {str(e)}"
        }), 500


@cotizaciones_bp.route("/plantilla")
def plantilla():
    """Descarga plantilla Excel para cotizaciones"""
    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Cotizaciones"
        ws["A1"] = "Fecha"
        ws["B1"] = "Cotización DL"
        ws["A2"] = date.today().strftime("%d/%m/%Y")
        ws["B2"] = "7,61982"
        for cell in ws["A"]:
            cell.number_format = "DD/MM/YYYY"
        ws.column_dimensions["A"].width = 14
        ws.column_dimensions["B"].width = 18
        
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
            wb.save(tmp.name)
            tmp_path = tmp.name
            
        from flask import send_file
        return send_file(tmp_path, as_attachment=True, download_name="plantilla_cotizaciones.xlsx")
        
    except Exception as e:
        logger.error(f"Error generando plantilla: {e}")
        return jsonify({"error": str(e)}), 500


@cotizaciones_bp.route("/importar_excel", methods=["POST"])
def importar_excel():
    """Importa cotizaciones desde un archivo Excel"""
    try:
        data = request.json
        
        if not data or not data.get("file_b64"):
            return jsonify({
                "ok": False,
                "error": "No se recibió archivo"
            }), 400
            
        file_bytes = base64.b64decode(data["file_b64"])
        bases = data.get("bases", PAISES_PERMITIDOS)
        
        bases_validas = [b for b in bases if b in PAISES_PERMITIDOS]
        if not bases_validas:
            return jsonify({
                "ok": False,
                "error": "No hay bases válidas para importar"
            }), 400
        
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
            tmp.write(file_bytes)
            tmp_path = tmp.name
            
        items = []
        try:
            try:
                wb = openpyxl.load_workbook(tmp_path, data_only=True)
            except (openpyxl.utils.exceptions.InvalidFileException, zipfile.BadZipFile) as e:
                return jsonify({
                    "ok": False,
                    "error": "El archivo no es un Excel válido"
                }), 400
                
            ws = wb.active
            
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row[0] or not row[1]:
                    continue
                    
                # Validar fecha
                fecha_val = row[0]
                if hasattr(fecha_val, "strftime"):
                    fecha_str = fecha_val.strftime("%Y-%m-%d")
                else:
                    parts = str(fecha_val).split("/")
                    if len(parts) == 3:
                        try:
                            fecha_str = f"{parts[2]}-{parts[1].zfill(2)}-{parts[0].zfill(2)}"
                            if not validar_fecha(fecha_str):
                                continue
                        except:
                            continue
                    else:
                        continue
                
                # Validar cotización
                try:
                    coti_val = float(str(row[1]).replace(",", ".").replace(" ", ""))
                    if coti_val <= 0:
                        continue
                except (ValueError, TypeError):
                    continue
                    
                items.append({
                    "fecha": fecha_str,
                    "cotizacion": coti_val,
                    "bases": bases_validas
                })
                
        finally:
            os.unlink(tmp_path)
        
        if not items:
            return jsonify({
                "ok": False,
                "error": "No se encontraron datos válidos en el Excel."
            }), 400

        resultado = _guardar_cotizaciones(items)
        
        if resultado["bases_con_error"]:
            return jsonify({
                "ok": resultado["ok"],
                "insertados": resultado["insertados"],
                "errores": resultado["errores"],
                "bases_con_error": resultado["bases_con_error"],
                "mensaje": f"⚠️ No se pudo importar en: {', '.join(resultado['bases_con_error'])}."
            })

        return jsonify({
            "ok": resultado["ok"],
            "insertados": resultado["insertados"],
            "errores": resultado["errores"]
        })
        
    except Exception as e:
        logger.error(f"Error importando Excel: {e}")
        return jsonify({
            "ok": False,
            "error": f"Error inesperado: {str(e)}"
        }), 500


@cotizaciones_bp.route("/cotizacion")
def cotizacion():
    """
    Obtiene la cotización para una fecha y moneda específica.
    Intenta obtenerla de países prioritarios (primero RD) antes de devolver 1.0.
    """
    try:
        moneda = request.args.get("moneda", "DL")
        fecha = request.args.get("fecha", "")
        
        if moneda not in ['DL', 'PS']:
            return jsonify({
                "error": "Moneda inválida. Solo se permiten DL y PS",
                "code": "INVALID_CURRENCY"
            }), 400
        
        if moneda == "PS":
            return jsonify({"cotizacion": 1.0})
        
        if fecha and not validar_fecha(fecha):
            return jsonify({
                "error": "Fecha inválida. Formato esperado: YYYY-MM-DD",
                "code": "INVALID_DATE"
            }), 400

        # Intentar con países prioritarios (primero RD, luego el resto)
        for sigla in PAISES_PRIORITARIOS:
            conn = None
            try:
                conn = obtener_conexion_pais_segura(sigla)
                cursor = conn.cursor()

                query = """
                    SELECT TOP 1 COTI_COTIZACION as coti
                    FROM SIST_COTI
                    WHERE COTI_MONEDA1 = ? AND COTI_MONEDA2 = ?
                """
                params = [moneda, 'PS']

                if fecha:
                    query += " AND COTI_FECHA <= ?"
                    params.append(fecha)

                query += " ORDER BY COTI_FECHA DESC"

                cursor.execute(query, params)
                row = cursor.fetchone()
                
                if row and row[0]:
                    try:
                        cotizacion = float(str(row[0]).replace(",", "."))
                        logger.info(f"Cotización obtenida de {sigla}: {cotizacion}")
                        return jsonify({"cotizacion": cotizacion})
                    except (ValueError, TypeError):
                        continue
                        
            except pyodbc.Error as e:
                logger.warning(f"Error al consultar cotización en {sigla}: {e}")
            except Exception as e:
                logger.warning(f"Error inesperado en {sigla}: {e}")
            finally:
                if conn:
                    try:
                        conn.close()
                    except:
                        pass
        
        # Si ningún país devolvió cotización, retornar 1.0
        logger.warning("No se encontró cotización en ningún país, usando 1.0")
        return jsonify({"cotizacion": 1.0})
        
    except Exception as e:
        logger.error(f"Error en cotización: {e}")
        return jsonify({"cotizacion": 1.0})